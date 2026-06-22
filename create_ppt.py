from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from PIL import Image, ImageDraw
import base64, io, re

def try_decode_and_save(b64_data, path, size=200):
    try:
        d = re.sub(r'[^A-Za-z0-9+/=]', '', b64_data)
        d = d.rstrip('=')
        d += '=' * ((4 - len(d) % 4) % 4)
        img_bytes = base64.b64decode(d)
        img = Image.open(io.BytesIO(img_bytes)).convert("RGBA")
        img = img.resize((size, size), Image.LANCZOS)
    except:
        img = Image.new("RGBA", (size, size), (220, 200, 170, 255))

    mask = Image.new("L", (size, size), 0)
    ImageDraw.Draw(mask).ellipse((0, 0, size, size), fill=255)
    output = Image.new("RGBA", (size, size), (255, 255, 255, 0))
    output.paste(img, (0, 0), mask)
    border = Image.new("RGBA", (size, size), (0, 0, 0, 0))
    ImageDraw.Draw(border).ellipse((0, 0, size-1, size-1), outline=(224, 48, 48, 255), width=6)
    output = Image.alpha_composite(output, border)
    output.save(path, format="PNG")
    return path

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = RGBColor(0xFE, 0xF6, 0xE4)

tx = slide.shapes.add_textbox(Inches(1.5), Inches(1.2), Inches(10.333), Inches(0.8))
tx.text_frame.word_wrap = True
p = tx.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "신세계 강남점  ·  상반기 정년퇴임 기념"
r.font.size = Pt(36); r.font.bold = True; r.font.color.rgb = RGBColor(0x7B, 0x40, 0x20)

ln = slide.shapes.add_shape(1, Inches(3), Inches(2.15), Inches(7.333), Pt(1))
ln.fill.solid(); ln.fill.fore_color.rgb = RGBColor(0xC9, 0x85, 0x3A); ln.line.fill.background()

tx2 = slide.shapes.add_textbox(Inches(6.2), Inches(1.95), Inches(0.9), Inches(0.4))
p2 = tx2.text_frame.paragraphs[0]; p2.alignment = PP_ALIGN.CENTER
r2 = p2.add_run(); r2.text = "✿"; r2.font.size = Pt(20); r2.font.color.rgb = RGBColor(0xC9, 0x85, 0x3A)

tx3 = slide.shapes.add_textbox(Inches(1.5), Inches(2.5), Inches(10.333), Inches(1.2))
tx3.text_frame.word_wrap = True
p3 = tx3.text_frame.paragraphs[0]; p3.alignment = PP_ALIGN.CENTER; p3.space_after = Pt(8)
r3 = p3.add_run(); r3.text = "그동안 함께해 주셔서 정말 고마웠습니다"
r3.font.size = Pt(24); r3.font.color.rgb = RGBColor(0x6B, 0x4A, 0x28)
p4 = tx3.text_frame.add_paragraph(); p4.alignment = PP_ALIGN.CENTER
r4 = p4.add_run(); r4.text = "감사한 마음을 담아 특식을 준비했습니다  ·  맛있게 드세요!"
r4.font.size = Pt(24); r4.font.color.rgb = RGBColor(0x6B, 0x4A, 0x28)

sb = slide.shapes.add_shape(1, Inches(0), Inches(4.2), Inches(13.333), Inches(3.3))
sb.fill.solid(); sb.fill.fore_color.rgb = RGBColor(0xFA, 0xE8, 0xC2); sb.line.fill.background()

people = [
    {"name": "김보영", "team": "POP운영팀", "years": "23년 1개월"},
    {"name": "전은영", "team": "POP운영팀", "years": "16년 8개월"},
    {"name": "장정원", "team": "POP운영팀", "years": "16년 1개월"},
    {"name": "제경순", "team": "POP운영팀", "years": "15년 8개월"},
    {"name": "최복례", "team": "식품팀", "years": "13년 1개월"},
]

with open("/home/user/claude/retirement-banner.html", "r") as f:
    html = f.read()
b64_images = re.findall(r'src="data:image/jpeg;base64,([^"]+)"', html)

start_x, gap, ps = 1.5, 2.2, 1.5

for i, person in enumerate(people):
    x = start_x + i * gap

    tb = slide.shapes.add_textbox(Inches(x - 0.15), Inches(4.4), Inches(ps + 0.3), Inches(0.35))
    pp = tb.text_frame.paragraphs[0]; pp.alignment = PP_ALIGN.CENTER
    rr = pp.add_run(); rr.text = person["team"]
    rr.font.size = Pt(12); rr.font.bold = True; rr.font.color.rgb = RGBColor(0xB8, 0x70, 0x20)

    img_path = f"/tmp/person_{i}.png"
    if i < len(b64_images):
        try_decode_and_save(b64_images[i], img_path)
    slide.shapes.add_picture(img_path, Inches(x), Inches(4.85), Inches(ps), Inches(ps))

    for text, y, sz, bold, color in [
        (person["name"], 6.4, 18, True, (0x2C, 0x15, 0x08)),
        ("파트너님", 6.7, 13, False, (0x9B, 0x73, 0x40)),
        (person["years"], 6.95, 12, True, (0xC9, 0x85, 0x3A)),
    ]:
        tb = slide.shapes.add_textbox(Inches(x - 0.15), Inches(y), Inches(ps + 0.3), Inches(0.3))
        pp = tb.text_frame.paragraphs[0]; pp.alignment = PP_ALIGN.CENTER
        rr = pp.add_run(); rr.text = text
        rr.font.size = Pt(sz); rr.font.bold = bold; rr.font.color.rgb = RGBColor(*color)

prs.save("/home/user/claude/retirement-banner.pptx")
print("PPT created successfully!")
